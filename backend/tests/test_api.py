import io

from openpyxl import load_workbook

from app.services.demo import DEMO_CSV
from tests.conftest import drain


async def _upload(client, content: bytes, name: str = "goods.csv"):
    return await client.post("/api/uploads/preview", files={"file": (name, content, "text/csv")})


async def _create(client, preview: dict, **settings):
    body = {
        "filename": preview["filename"],
        "mapping": preview["suggested_mapping"],
        "rows": preview["rows"],
        "marketplace": "wb",
        "tone": "neutral",
        "length": "short",
        **settings,
    }
    return await client.post("/api/batches", json=body)


async def test_demo_batch_is_ready_on_first_visit(client):
    resp = await client.get("/api/batches")
    assert resp.status_code == 200
    [demo] = resp.json()
    assert demo["is_demo"] and demo["total"] == 30
    await drain()
    batch = (await client.get(f"/api/batches/{demo['id']}")).json()
    assert batch["counts"]["done"] == 30


async def test_demo_photos_are_measured_in_advance(client):
    [demo] = (await client.get("/api/batches")).json()
    batch = (await client.get(f"/api/batches/{demo['id']}")).json()
    for p in batch["products"]:
        assert p["source"]["photo_width"] and p["source"]["photo_height"], p["source"]["sku"]
        # Все фото демо подходят площадкам — замечаний к ним нет
        assert not [w for w in p["warnings"] if w["code"] == "photo"], p["source"]["sku"]


async def test_old_demo_gets_photo_sizes_on_open(client):
    from sqlalchemy import select

    from app.db import SessionLocal
    from app.models import Product

    [demo] = (await client.get("/api/batches")).json()
    # Демо, созданное до того, как размеры фото стали хранить
    async with SessionLocal() as session:
        for product in await session.scalars(select(Product).where(Product.batch_id == demo["id"])):
            product.source = {**product.source, "photo_width": None, "photo_height": None}
        await session.commit()
    batch = (await client.get(f"/api/batches/{demo['id']}")).json()
    assert all(p["source"]["photo_width"] for p in batch["products"])


async def test_visitors_do_not_see_each_other(client):
    await client.get("/api/batches")
    other = await client.get("/api/batches", headers={"X-Visitor-Id": "someone-else-02"})
    mine = await client.get("/api/batches")
    assert other.json()[0]["id"] != mine.json()[0]["id"]
    resp = await client.get(
        f"/api/batches/{mine.json()[0]['id']}", headers={"X-Visitor-Id": "someone-else-02"}
    )
    assert resp.status_code == 404


async def test_full_flow_upload_generate_edit_export(client):
    preview = (await _upload(client, DEMO_CSV.read_bytes())).json()
    assert preview["suggested_mapping"]["name"] == "Название"
    assert preview["suggested_mapping"]["attributes"] == "Характеристики"
    assert len(preview["rows"]) == 30

    resp = await _create(client, preview)
    assert resp.status_code == 201, resp.text
    batch = resp.json()
    await drain()
    batch = (await client.get(f"/api/batches/{batch['id']}")).json()
    assert batch["counts"]["done"] == 30

    product = batch["products"][0]
    assert len(product["result"]["title"]) <= 60  # лимит WB

    resp = await client.patch(
        f"/api/products/{product['id']}", json={"title": "Плед " + "очень длинный " * 6}
    )
    edited = resp.json()
    assert edited["edited"]
    assert any(w["code"] == "too_long" and w["field"] == "title" for w in edited["warnings"])

    resp = await client.post(
        f"/api/products/{product['id']}/regenerate", json={"field": "title", "comment": "короче"}
    )
    assert resp.status_code == 200
    assert len(resp.json()["result"]["title"]) <= 60

    resp = await client.get(f"/api/batches/{batch['id']}/export")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["Карточки"]
    assert ws.cell(1, 1).value == "Артикул продавца"
    assert ws.max_row == 31
    assert ws.cell(2, 2).value == "Текстиль"  # кириллица на месте


async def test_retry_failed_rows(client):
    preview = (await _upload(client, DEMO_CSV.read_bytes())).json()
    batch = (await _create(client, preview)).json()
    await drain()

    from app.db import SessionLocal
    from app.models import Product

    async with SessionLocal() as s:
        p = await s.get(Product, batch["products"][3]["id"])
        p.status, p.error = "error", "Таймаут модели"
        await s.commit()

    resp = await client.post(f"/api/batches/{batch['id']}/retry-failed")
    assert resp.json()["counts"]["error"] == 0
    await drain()
    batch = (await client.get(f"/api/batches/{batch['id']}")).json()
    assert batch["counts"]["done"] == 30


async def test_limits(client):
    many = "Артикул,Название\n" + "\n".join(f"A{i},Товар {i}" for i in range(51))
    resp = await _upload(client, many.encode())
    assert resp.status_code == 422
    assert "50" in resp.json()["detail"]

    preview = (await _upload(client, "Артикул,Название\nA1,Кружка\n".encode())).json()
    for _ in range(3):
        assert (await _create(client, preview)).status_code == 201
    resp = await _create(client, preview)
    assert resp.status_code == 429
    assert "Retry-After" in resp.headers


async def test_demo_reset_restores_original(client):
    [demo] = (await client.get("/api/batches")).json()
    await drain()
    batch = (await client.get(f"/api/batches/{demo['id']}")).json()
    pid = batch["products"][0]["id"]
    await client.patch(f"/api/products/{pid}", json={"title": "Испорчено"})

    resp = await client.post("/api/demo/reset")
    assert resp.status_code == 200
    await drain()
    batches = (await client.get("/api/batches")).json()
    assert len(batches) == 1 and batches[0]["id"] != demo["id"]


async def test_demo_is_not_recreated_after_deleting_everything(client):
    [demo] = (await client.get("/api/batches")).json()
    assert (await client.delete(f"/api/batches/{demo['id']}")).status_code == 204
    # Повторный визит: посетитель уже известен, демо само не возвращается
    assert (await client.get("/api/batches")).json() == []
    # …но возвращается кнопкой «Сбросить демо»
    resp = await client.post("/api/demo/reset")
    assert resp.status_code == 200 and resp.json()["is_demo"]
    assert len((await client.get("/api/batches")).json()) == 1


async def test_existing_visitor_without_record_gets_no_duplicate_demo(client):
    """Посетители, пришедшие до появления таблицы visitors, не получают второе демо."""
    from app.db import SessionLocal
    from app.models import Visitor

    await client.get("/api/batches")
    async with SessionLocal() as s:
        await s.delete(await s.get(Visitor, "test-visitor-0001"))
        await s.commit()
    assert len((await client.get("/api/batches")).json()) == 1


async def test_cleanup_forgets_old_visitors(client):
    from datetime import UTC, datetime, timedelta

    from sqlalchemy import update

    from app.db import SessionLocal
    from app.models import Batch, Visitor
    from app.services.demo import cleanup_old

    [demo] = (await client.get("/api/batches")).json()
    old = datetime.now(UTC) - timedelta(days=2)
    async with SessionLocal() as s:
        await s.execute(update(Batch).values(created_at=old))
        await s.execute(update(Visitor).values(created_at=old))
        await s.commit()
        assert await cleanup_old(s, ttl_hours=24) == 1
    # На следующий день посетитель снова видит свежую демо-партию
    [fresh] = (await client.get("/api/batches")).json()
    assert fresh["is_demo"] and fresh["id"] != demo["id"]


async def test_restore_demo_keeps_other_batches(client):
    [demo] = (await client.get("/api/batches")).json()
    preview = (await _upload(client, "Артикул,Название\nA1,Кружка\n".encode())).json()
    own = (await _create(client, preview)).json()
    await client.delete(f"/api/batches/{demo['id']}")

    restored = (await client.post("/api/demo/restore")).json()
    assert restored["is_demo"]
    ids = {b["id"] for b in (await client.get("/api/batches")).json()}
    assert ids == {own["id"], restored["id"]}
    # Повторный вызов не плодит демо
    assert (await client.post("/api/demo/restore")).json()["id"] == restored["id"]
