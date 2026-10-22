"""Разбор CSV/XLSX и сопоставление колонок с полями сервиса."""

import csv
import io
import re

from openpyxl import load_workbook

from app.schemas import Attribute, SourceField, SourceProduct


class ParseError(ValueError):
    pass


SYNONYMS: dict[SourceField, list[str]] = {
    "sku": ["артикул", "sku", "код", "код товара", "vendor code", "артикул продавца", "id"],
    "name": ["название", "наименование", "name", "товар", "title", "наименование товара"],
    "category": ["категория", "category", "тип", "предмет", "раздел", "тип товара"],
    "brand": ["бренд", "brand", "производитель", "марка", "торговая марка"],
    "attributes": ["характеристики", "attributes", "свойства", "параметры", "атрибуты", "specs"],
    "photo": ["фото", "photo", "изображение", "image", "картинка", "ссылка на фото", "url фото"],
}


def _norm(header: str) -> str:
    return re.sub(r"[\s_*]+", " ", header.lower()).strip(" :")


def suggest_mapping(columns: list[str]) -> dict[SourceField, str | None]:
    mapping: dict[SourceField, str | None] = {}
    used: set[str] = set()
    for field, names in SYNONYMS.items():
        mapping[field] = None
        for col in columns:
            if col not in used and _norm(col) in names:
                mapping[field] = col
                used.add(col)
                break
        if mapping[field] is None:
            # «Название товара (рус)» тоже подходит под «название»
            for col in columns:
                if col not in used and any(_norm(col).startswith(n) for n in names):
                    mapping[field] = col
                    used.add(col)
                    break
    return mapping


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParseError("Не удалось определить кодировку файла — сохраните его в UTF-8")


def _read_csv(data: bytes) -> list[list[str]]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(data: bytes) -> list[list[str]]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # openpyxl бросает разные исключения на битых файлах
        raise ParseError("Файл не похож на XLSX") from e
    ws = wb.worksheets[0]
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if v is None else str(v) for v in row])
    wb.close()
    return rows


def parse_table(
    filename: str, data: bytes, max_rows: int, max_cell: int
) -> tuple[list[str], list[dict[str, str]]]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        raw = _read_csv(data)
    elif lower.endswith((".xlsx", ".xlsm")):
        raw = _read_xlsx(data)
    else:
        raise ParseError("Поддерживаются только CSV и XLSX")

    raw = [[c.strip() for c in row] for row in raw if any(c.strip() for c in row)]
    if len(raw) < 2:
        raise ParseError("В файле нет строк с товарами")

    header = raw[0]
    columns: list[str] = []
    for i, col in enumerate(header):
        name = col or f"Колонка {i + 1}"
        while name in columns:
            name += " (2)"
        columns.append(name)

    body = raw[1:]
    if len(body) > max_rows:
        raise ParseError(f"В файле {len(body)} строк, в демо можно не больше {max_rows}")

    rows = []
    for row in body:
        rows.append(
            {col: (row[i] if i < len(row) else "")[:max_cell] for i, col in enumerate(columns)}
        )
    return columns, rows


def parse_attributes(text: str) -> list[Attribute]:
    """«Материал: хлопок; Размер: 200x220 см» → пары ключ — значение."""
    attrs: list[Attribute] = []
    extras: list[str] = []
    for part in re.split(r"[;\n|]+", text):
        part = part.strip()
        if not part:
            continue
        m = re.match(r"^([^:=—–]{1,60})\s*[:=—–]\s*(.+)$", part)
        if m:
            attrs.append(Attribute(name=m.group(1).strip(), value=m.group(2).strip()))
        else:
            extras.append(part)
    if extras:
        attrs.append(Attribute(name="Особенности", value=", ".join(extras)))
    return attrs


def build_source(
    row: dict[str, str], mapping: dict[SourceField, str | None], extra_as_attributes: bool
) -> SourceProduct:
    def get(field: SourceField) -> str:
        col = mapping.get(field)
        return row.get(col, "").strip() if col else ""

    attributes = parse_attributes(get("attributes"))
    if extra_as_attributes:
        mapped = {c for c in mapping.values() if c}
        known = {a.name.lower() for a in attributes}
        for col, value in row.items():
            if col not in mapped and value.strip() and col.lower() not in known:
                attributes.append(Attribute(name=col, value=value.strip()))

    return SourceProduct(
        sku=get("sku"),
        name=get("name"),
        category=get("category"),
        brand=get("brand"),
        attributes=attributes,
        photo=get("photo"),
    )
