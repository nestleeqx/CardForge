"""Сборка XLSX с колонками под шаблон выбранного маркетплейса."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import Batch
from app.schemas import CardResult, SourceProduct

# (заголовок, ширина) для базовых колонок; характеристики добавляются справа
TEMPLATES = {
    "ozon": [
        ("Артикул*", 16),
        ("Название товара", 50),
        ("Бренд*", 18),
        ("Тип*", 22),
        ("Аннотация", 80),
        ("Ключевые слова", 50),
        ("Ссылка на главное фото*", 30),
    ],
    "wb": [
        ("Артикул продавца", 16),
        ("Предмет", 22),
        ("Бренд", 18),
        ("Наименование", 40),
        ("Описание", 80),
        ("Ключевые слова", 50),
        ("Фото", 30),
    ],
}

HEADER_FILL = PatternFill("solid", fgColor="E8EEF9")
WARN_FILL = PatternFill("solid", fgColor="FFF4D6")


def _base_row(marketplace: str, src: SourceProduct, card: CardResult) -> list[str]:
    keywords = ", ".join(card.keywords)
    if marketplace == "ozon":
        return [src.sku, card.title, src.brand, src.category, card.description, keywords, src.photo]
    return [src.sku, src.category, src.brand, card.title, card.description, keywords, src.photo]


def build_xlsx(batch: Batch) -> bytes:
    template = TEMPLATES[batch.marketplace]
    done = [p for p in batch.products if p.status == "done" and p.result]

    attr_names: list[str] = []
    for p in done:
        for a in CardResult.model_validate(p.result).attributes:
            if a.name not in attr_names:
                attr_names.append(a.name)

    wb = Workbook()
    ws = wb.active
    ws.title = "Карточки"
    headers = [h for h, _ in template] + attr_names
    ws.append(headers)
    for i, (_, width) in enumerate(template, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for i in range(len(template) + 1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "B2"

    for p in done:
        src = SourceProduct.model_validate(p.source)
        card = CardResult.model_validate(p.result)
        values = {a.name: a.value for a in card.attributes}
        ws.append(_base_row(batch.marketplace, src, card) + [values.get(n, "") for n in attr_names])
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    checks = wb.create_sheet("Проверки")
    checks.append(["Артикул", "Статус", "Поле", "Предупреждение"])
    for cell in checks[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for col, width in zip("ABCD", (16, 12, 16, 90), strict=True):
        checks.column_dimensions[col].width = width
    for p in batch.products:
        sku = p.source.get("sku", "")
        if p.status == "error":
            checks.append([sku, "ошибка", "", p.error or "Не удалось сгенерировать"])
        for w in p.warnings:
            checks.append([sku, "готово", w.get("field") or "", w["message"]])
            for cell in checks[checks.max_row]:
                cell.fill = WARN_FILL

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
